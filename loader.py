# -*- coding: utf-8 -*-
"""文档加载与清洗：支持 .txt / .md / .pdf / .docx

这是 RAG 流水线的第一步，也是企业知识库最脏最累的一步。
"""
import re
from pathlib import Path

SUPPORTED = {".txt", ".md", ".pdf", ".docx"}


def _load_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError:
        raise RuntimeError("读 PDF 需要 pypdf：pip install pypdf")
    reader = PdfReader(str(path))
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def _load_docx(path: Path) -> str:
    try:
        from docx import Document
    except ImportError:
        raise RuntimeError("读 Word 需要 python-docx：pip install python-docx")
    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:                       # 表格也要收进来
        for row in table.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(parts)


def _load_text(path: Path) -> str:
    for enc in ("utf-8", "gbk", "utf-8-sig"):
        try:
            return path.read_text(encoding=enc)
        except UnicodeDecodeError:
            continue
    return path.read_text(encoding="utf-8", errors="ignore")


def clean(text: str) -> str:
    """清洗：这一步做得好不好，直接决定检索命中率。"""
    text = text.replace("\u3000", " ")            # 全角空格
    text = re.sub(r"[ \t]+", " ", text)            # 压缩连续空白
    text = re.sub(r"\n{3,}", "\n\n", text)         # 压缩连续空行
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", text)  # 控制字符
    # 去掉页眉页脚这类整页重复的噪声（按行判断）
    lines = [ln.strip() for ln in text.split("\n")]
    lines = [ln for ln in lines if not re.fullmatch(r"[-—=·\.\s]*第?\s*\d+\s*页?[-—=·\.\s]*", ln or "x")]
    return "\n".join(lines).strip()


def load_document(path: Path):
    """返回 (文本, 元数据)"""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED:
        raise ValueError(f"不支持的格式：{suffix}")
    if suffix == ".pdf":
        raw = _load_pdf(path)
    elif suffix == ".docx":
        raw = _load_docx(path)
    else:
        raw = _load_text(path)
    meta = {"source": path.name, "path": str(path), "chars": len(raw)}
    return clean(raw), meta


def load_dir(docs_dir: Path):
    """批量加载目录下所有支持的文档"""
    docs_dir = Path(docs_dir)
    out = []
    for p in sorted(docs_dir.rglob("*")):
        if p.is_file() and p.suffix.lower() in SUPPORTED:
            try:
                text, meta = load_document(p)
                if text:
                    out.append((text, meta))
            except Exception as e:
                print(f"  [跳过] {p.name}: {e}")
    return out


if __name__ == "__main__":
    from config import DOCS_DIR
    for text, meta in load_dir(DOCS_DIR):
        print(f"{meta['source']:24s} {meta['chars']:>7d} 字符")
